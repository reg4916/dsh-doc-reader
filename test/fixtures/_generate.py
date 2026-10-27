#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""One-shot generator for the dsh-doc-reader fixture corpus.

Runs inside test/fixtures only. Deleted after the corpus is verified.
PDF is hand-crafted (uncompressed, Helvetica, 3 pages) so the page count and
the literal PAGE-n-MARKER strings are deterministic for any extractor.
"""
import json
import random
import re
from pathlib import Path

FX = Path(__file__).resolve().parent
FENCE = chr(96) * 3

# ---------------------------------------------------------------- 1. txt ---
(FX / "sample.txt").write_text(
    "Hello from the dsh-doc-reader text fixture.\n"
    "\u041f\u0440\u0438\u0432\u0435\u0442, \u043c\u0438\u0440!\n"
    "This file is UTF-8 and mixes Latin and Cyrillic on purpose.\n"
    "\u041f\u0440\u043e\u0432\u0435\u0440\u043a\u0430: \u043a\u0438\u0440\u0438\u043b\u043b\u0438\u0446\u0430, \u0442\u0438\u0440\u0435 \u2014 \u0438 \u043a\u043e\u0435-\u0447\u0442\u043e \u0435\u0449\u0451.\n",
    encoding="utf-8",
)

# ----------------------------------------------------------------- 2. md ---
md_lines = [
    "# Sample Markdown Document",
    "",
    "\u041a\u0440\u0430\u0442\u043a\u043e\u0435 \u043e\u043f\u0438\u0441\u0430\u043d\u0438\u0435 \u0444\u0438\u043a\u0441\u0442\u0443\u0440\u044b \u0434\u043b\u044f \u043f\u043b\u0430\u0433\u0438\u043d\u0430 dsh-doc-reader.",
    "",
    "## \u0421\u043f\u0438\u0441\u043e\u043a",
    "",
    "- \u041f\u0435\u0440\u0432\u044b\u0439 \u043f\u0443\u043d\u043a\u0442 \u0441\u043f\u0438\u0441\u043a\u0430",
    "- \u0412\u0442\u043e\u0440\u043e\u0439 \u043f\u0443\u043d\u043a\u0442 \u0441\u043f\u0438\u0441\u043a\u0430",
    "- Third item in English",
    "",
    "## \u041a\u043e\u0434",
    "",
    FENCE + "python",
    "def greet(name: str) -> str:",
    '    """Return a greeting for name."""',
    '    return f"Hello, {name}!"',
    FENCE,
    "",
]
(FX / "sample.md").write_text("\n".join(md_lines), encoding="utf-8")

# ---------------------------------------------------------------- 3. csv ---
(FX / "sample.csv").write_text(
    "\u0413\u043e\u0434,\u0412\u044b\u0440\u0443\u0447\u043a\u0430\n2024,1000\n2025,2000\n2026,3000\n",
    encoding="utf-8",
)

# --------------------------------------------------------------- 4. docx ---
from docx import Document

doc = Document()
doc.core_properties.title = "DOCX fixture"
doc.core_properties.author = "dsh-doc-reader fixtures"

doc.add_heading("DOCX-TITLE-MARKER", level=0)
doc.add_heading("Section One: fixture overview", level=1)
doc.add_paragraph(
    "\u041a\u0438\u0440\u0438\u043b\u043b\u0438\u0447\u0435\u0441\u043a\u0438\u0439 \u0430\u0431\u0437\u0430\u0446: "
    "\u0414\u041e\u041a\u0423\u041c\u0415\u041d\u0422-\u041c\u0410\u0420\u041a\u0415\u0420 \u2014 "
    "\u043f\u0440\u043e\u0432\u0435\u0440\u043a\u0430 UTF-8 \u0432 DOCX."
)
doc.add_paragraph("First bullet item", style="List Bullet")
doc.add_paragraph("\u0412\u0442\u043e\u0440\u043e\u0439 \u043f\u0443\u043d\u043a\u0442 \u0441\u043f\u0438\u0441\u043a\u0430", style="List Bullet")

table = doc.add_table(rows=2, cols=3)
table.style = "Table Grid"
cells = [["A1", "B1", "C1"], ["A2", "B2", "C2"]]
for r, row in enumerate(cells):
    for c, value in enumerate(row):
        table.cell(r, c).text = value

doc.save(str(FX / "sample.docx"))

# --------------------------------------------------------------- 5. xlsx ---
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "Data"
ws.append(["Year", "Revenue", "Note"])
ws.append([2024, 1000, "ok"])
ws.append([2025, 2000, "\u041f\u0440\u0438\u0432\u0435\u0442"])

ws2 = wb.create_sheet("Second")
ws2["A1"] = "MERGED-MARKER"
ws2.merge_cells("A1:B1")
ws2["A2"] = 5
ws2["A3"] = 7
ws2["A4"] = "=SUM(A2:A3)"
wb.save(str(FX / "sample.xlsx"))

# --------------------------------------------------------------- 6. pptx ---
from pptx import Presentation

prs = Presentation()
Layout = prs.slide_layouts[1]  # Title and Content


def add_slide(title, bullets):
    slide = prs.slides.add_slide(Layout)
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.text = bullets[0]
    for bullet in bullets[1:]:
        body.add_paragraph().text = bullet
    return slide


s1 = add_slide("SLIDE-1-TITLE", ["BULLET-1-A", "BULLET-1-B"])
notes = s1.notes_slide.notes_text_frame
notes.text = "NOTES-MARKER: \u0417\u0430\u043c\u0435\u0442\u043a\u0438 \u0434\u043e\u043a\u043b\u0430\u0434\u0447\u0438\u043a\u0430 \u0434\u043b\u044f \u043f\u0435\u0440\u0432\u043e\u0433\u043e \u0441\u043b\u0430\u0439\u0434\u0430."
add_slide("SLIDE-2-TITLE", ["BULLET-2-A", "BULLET-2-B"])
add_slide("SLIDE-3-TITLE", ["BULLET-3-A"])
prs.save(str(FX / "sample.pptx"))

# ---------------------------------------------------------------- 7. pdf ---
def _escape(text):
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def _content_stream(lines):
    parts = ["BT", "/F1 16 Tf", "72 720 Td", "(%s) Tj" % _escape(lines[0]), "/F1 12 Tf", "0 -30 Td"]
    for line in lines[1:]:
        parts.append("(%s) Tj" % _escape(line))
        parts.append("0 -22 Td")
    parts.append("ET")
    return ("\n".join(parts) + "\n").encode("ascii")


PDF_PAGES = [
    [
        "PAGE-1-MARKER",
        "This is page one of the dsh-doc-reader PDF fixture.",
        "The document contains exactly three pages.",
        "Each page carries its own unique ASCII marker so tests can",
        "prove that extraction walks the whole document.",
        "Content streams are stored uncompressed for portability.",
    ],
    [
        "PAGE-2-MARKER",
        "Page two continues with ordinary descriptive prose.",
        "Automated tests count the pages and check that every marker",
        "is present in the extracted text.",
        "There is no scanned image anywhere in this file.",
    ],
    [
        "PAGE-3-MARKER",
        "Page three is the final page of the fixture.",
        "Total page count: three.",
        "End of document.",
    ],
]


def build_pdf(pages):
    n = len(pages)
    pairs = []
    next_id = 4
    for _ in pages:
        pairs.append((next_id, next_id + 1))
        next_id += 2
    max_id = next_id - 1
    objs = {}
    objs[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join("%d 0 R" % pid for pid, _ in pairs)
    objs[2] = ("<< /Type /Pages /Count %d /Kids [%s] >>" % (n, kids)).encode("ascii")
    objs[3] = (b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
               b"/Encoding /WinAnsiEncoding >>")
    for (pid, cid), lines in zip(pairs, pages):
        stream = _content_stream(lines)
        objs[pid] = (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            "/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % cid
        ).encode("ascii")
        objs[cid] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"endstream\n"

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = {}
    for i in range(1, max_id + 1):
        offsets[i] = len(out)
        out += b"%d 0 obj\n" % i + objs[i] + b"\nendobj\n"
    xref_at = len(out)
    out += b"xref\n0 %d\n" % (max_id + 1)
    out += b"0000000000 65535 f \n"
    for i in range(1, max_id + 1):
        out += b"%010d 00000 n \n" % offsets[i]
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (max_id + 1, xref_at)
    return bytes(out)


pdf_bytes = build_pdf(PDF_PAGES)
(FX / "sample.pdf").write_bytes(pdf_bytes)
assert len(re.findall(rb"/Type\s*/Page\b", pdf_bytes)) == 3, "pdf page objects != 3"
assert re.search(rb"/Count 3\b", pdf_bytes), "pdf /Count != 3"

# ------------------------------------------------------------ 8. edges ---
(FX / "empty.txt").write_bytes(b"")

rnd = random.Random(20261026)
(FX / "corrupt.pdf").write_bytes(bytes(rnd.randrange(256) for _ in range(1024)))

(FX / "not-a-docx.docx").write_text(
    "This file is plain UTF-8 text, not a DOCX (ZIP) container.\n"
    "NOT-A-DOCX-MARKER\n"
    "It exists to exercise the reader's error path for a misleading extension.\n",
    encoding="utf-8",
)

empty_wb = Workbook()
empty_wb.save(str(FX / "empty.xlsx"))

# --------------------------------------------------------- 9. expected ---
expected = {
    "sample.txt": {
        "format": "text",
        "substrings": ["Hello", "\u041f\u0440\u0438\u0432\u0435\u0442", "UTF-8"],
    },
    "sample.md": {
        "format": "text",
        "substrings": ["Sample Markdown Document", "\u041f\u0435\u0440\u0432\u044b\u0439 \u043f\u0443\u043d\u043a\u0442 \u0441\u043f\u0438\u0441\u043a\u0430", "def greet"],
    },
    "sample.csv": {
        "format": "text",
        "substrings": ["\u0413\u043e\u0434", "\u0412\u044b\u0440\u0443\u0447\u043a\u0430", "2026"],
    },
    "sample.docx": {
        "format": "docx",
        "substrings": ["DOCX-TITLE-MARKER", "\u0414\u041e\u041a\u0423\u041c\u0415\u041d\u0422-\u041c\u0410\u0420\u041a\u0415\u0420", "A1", "C2"],
    },
    "sample.xlsx": {
        "format": "xlsx",
        "sheets": ["Data", "Second"],
        "substrings": ["Year", "Revenue", "\u041f\u0440\u0438\u0432\u0435\u0442", "MERGED-MARKER"],
    },
    "sample.pptx": {
        "format": "pptx",
        "slides": 3,
        "substrings": ["SLIDE-1-TITLE", "SLIDE-3-TITLE", "NOTES-MARKER"],
    },
    "sample.pdf": {
        "format": "pdf",
        "pages": 3,
        "substrings": ["PAGE-1-MARKER", "PAGE-2-MARKER", "PAGE-3-MARKER"],
    },
    "empty.txt": {"format": "text", "empty": True, "substrings": []},
    "empty.xlsx": {"format": "xlsx", "empty": True, "sheets": ["Sheet"], "substrings": []},
    "not-a-docx.docx": {"format": "docx", "invalid": True, "substrings": []},
    "corrupt.pdf": {"format": "pdf", "invalid": True, "substrings": []},
}
(FX / "expected.json").write_text(
    json.dumps(expected, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
)

print("generated:", sorted(p.name for p in FX.iterdir()))
